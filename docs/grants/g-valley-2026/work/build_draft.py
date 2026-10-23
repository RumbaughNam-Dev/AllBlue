from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json

ROOT=Path(__file__).resolve().parents[1]
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8)
sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 s=doc.styles[name];s.font.name='SUIT';s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'SUIT')
 s.font.size=Pt(10.5)
 s.paragraph_format.space_after=Pt(7)
 s.paragraph_format.line_spacing=1.25
doc.styles['Title'].font.size=Pt(21)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
h=sec.header.paragraphs[0];h.text='AllBlue  |  G밸리 창업경진대회  |  사업계획서'
h.runs[0].font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('2026년 10월 7일 작성  ·  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
for r in f.runs:r.font.size=Pt(8)

def p(text,style=None):return doc.add_paragraph(text,style)
def h2(text):doc.add_heading(text,2)
def page(title):
 h=doc.add_heading(title,1);h.paragraph_format.page_break_before=True
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,txt in zip(t.rows[0].cells,headers):c.text=txt
 for row in rows:
  for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
 for ri,row in enumerate(t.rows):
  for ci,c in enumerate(row.cells):
   if widths:c.width=Cm(widths[ci])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   pr.append(b)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'95');x.set(qn('w:type'),'dxa');margins.append(x)
   pr.append(margins)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');pr.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.15
    for r in para.runs:r.font.size=Pt(9);r.bold=ri==0
  nr=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(nr)
 rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
 p('')
 return t

doc.add_heading('AllBlue 사업계획서',0)
p('2026 제10회 G밸리 창업경진대회  |  럼바우  |  대표 남상현')
doc.add_heading('1 문제인식',1)
h2('프리다이빙 수업의 준비와 기록을 한 곳에서')
p('AllBlue는 프리다이빙 강사가 일정과 수강생, 교육 진도, 동의서와 수업 피드백을 관리하는 모바일 앱입니다. 수업을 준비할 때 필요한 정보와 수업이 끝난 뒤 남기는 기록을 연결하는 데 초점을 맞췄습니다. 강사는 다음 수업을 준비하면서 수강생이 어디까지 배웠고 어떤 부분을 더 연습해야 하는지 확인할 수 있습니다.')
p('저는 개발 14년차이자 프리다이빙 3년차로, AIDA Level 2 자격을 보유하고 있습니다. 개발과 프리다이빙 경험을 함께 살릴 수 있는 사업으로 AllBlue를 준비하고 있습니다. 현재 운영 환경에서 앱을 테스트하고 있으며, 강사와 수강생을 합한 실제 사용자는 5명 이내입니다.')
h2('해결하려는 문제')
p('제가 파악한 프리다이빙 교육 현장에서는 엑셀로 수강생과 교육 진행 상황을 관리하는 경우가 많습니다. 엑셀은 항목을 자유롭게 정리하기에 익숙한 도구입니다. 다만 일정 안내와 동의서, 수업 후 피드백을 따로 관리하면 한 수강생의 정보를 여러 곳에서 찾아야 합니다. AllBlue는 엑셀에 정리하던 교육 진도를 실제 수업 일정과 연결하고, 수업 직후 휴대전화에서 기록을 남기도록 구성했습니다.')
p('AllBlue에서는 한 수업에 필요한 내용을 참가자별로 모아 보도록 구성했습니다. 강사가 앱을 계속 사용할 이유는 실제 수업을 준비하고 기록하는 일이 얼마나 편해지는지에 있다고 봅니다. 초기 사용 강사와 함께 기존 방식과 비교하며 준비 시간과 기록에 드는 부담을 확인할 계획입니다.')
h2('처음 집중할 고객')
table(['구분','대상과 이용 목적'],[
 ['비용을 지불하는 고객','반복적으로 교육과정을 운영하는 프리다이빙 강사'],
 ['함께 사용하는 이용자','일정을 확인하고 교육 진도와 피드백을 공유하는 수강생'],
 ['추후 검토할 고객','여러 강사가 함께 수업을 운영하는 소규모 교육업체']],[4,13])
p('초기에는 국내 프리다이빙 강사를 대상으로 사용성을 확인하겠습니다. 전체 다이빙 인구보다 실제로 수업을 운영하고 관리 도구에 비용을 지불할 강사의 수가 중요합니다. 협업 채널을 통해 접촉한 강사의 교육 빈도와 도입 의향을 파악하고, 그 결과를 바탕으로 모집 규모를 정할 계획입니다.')

page('2 해결방안')
h2('한 번의 수업을 준비하고 마무리하는 흐름')
p('강사가 일정을 만들고 수강생을 초대하면 수강생은 참석 여부를 응답합니다. 강사는 일정 상세에서 참가자별 서류 준비 상태와 교육 정보를 확인합니다. 수업 뒤에는 이수 항목과 디브리핑을 남기고, 다음 수업에서 다시 참고합니다. 디브리핑은 수업 중 잘된 점과 보완할 점을 정리하는 피드백 기록입니다.')
table(['단계','앱에 구성한 기능','사용 목적'],[
 ['수업 준비','일정 작성과 수정, 참가자 초대와 응답','누가 어떤 수업에 참여하는지 확인'],
 ['서류 확인','면책동의서와 의료진술서의 서명 상태 및 문서 조회','참가자의 준비 상태 확인'],
 ['교육 기록','교육 요구사항별 완료 상태 관리','수강생별 진행 상황 파악'],
 ['수업 마무리','디브리핑 작성과 이전 기록 조회','다음 수업에 필요한 피드백 보관'],
 ['관계 관리','강사와 수강생의 연결 관리','함께 교육하는 사람들의 정보 확인']],[2.7,7.4,6.9])
h2('제품 개발과 출시 준비')
p('Expo와 React Native를 사용해 앱을 개발했습니다. 일정, 참가자, 교육 진도와 디브리핑 데이터를 서버와 연결하는 구조입니다. 현재 App Store 검토는 완료됐고 Google Play는 검토 중입니다. 운영 환경에서 초대 응답부터 수업 후 기록까지 실제 사용 흐름을 점검하고 있습니다.')
p('출시 전에는 강사와 수강생이 각자 필요한 정보에 접근하는지, 문서 조회와 기록 저장이 제대로 동작하는지를 우선 확인하겠습니다. 특히 의료진술서 등 개인 정보가 담긴 문서는 수집 범위와 보관·삭제 절차를 점검하겠습니다.')
h2('경쟁 서비스와 비교할 부분')
p('다이빙 분야에는 MySSI처럼 디지털 교육, 자격증과 로그북을 제공하는 앱이 있습니다[1]. AllBlue는 강사가 매번 반복하는 수업 운영에 집중하려 합니다. 참석 여부와 서류를 확인하고 교육 진도와 피드백을 남기는 과정이 얼마나 자연스럽게 이어지는지를 살펴볼 계획입니다.')
p('비교 기준은 실제 강사가 수행하는 업무로 잡겠습니다. 기존에 쓰던 도구와 AllBlue에서 같은 수업을 준비해 보고, 정보를 찾는 시간과 기록 작성의 편의성을 비교하겠습니다. 초기 피드백에서 반복해서 나오는 불편부터 개선하겠습니다.')


source=p('[1] SSI, MySSI App. https://www.divessi.com/en/join-ssi/register-free/myssi-app (2026.10.7. 확인)')
source.paragraph_format.line_spacing=1.1
for run in source.runs:run.font.size=Pt(8)

page('2 교육 진도와 수업 피드백 화면')
p('엑셀로 정리하던 교육 항목을 수강생별 화면에서 확인하고, 수업이 끝난 자리에서 진행 상황과 피드백을 남길 수 있도록 구성했습니다. 아래는 AllBlue 앱 화면을 담은 iOS 소개 이미지입니다.')
fig=doc.add_table(rows=1,cols=2);fig.autofit=False
for cell,filename,caption in zip(fig.rows[0].cells,['05_feature_3.png','06_feature_4.png'],['교육 진도 확인','수업 후 디브리핑']):
 cell.width=Cm(8.5)
 para=cell.paragraphs[0];para.alignment=WD_ALIGN_PARAGRAPH.CENTER
 para.paragraph_format.space_after=Pt(4)
 run=para.add_run();pic=run.add_picture(str(ROOT/'sources/screenshots'/filename),width=Cm(7.7))
 pic._inline.docPr.set('descr',caption+'을 보여주는 AllBlue iOS 소개 이미지')
 para=cell.add_paragraph(caption);para.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for run in para.runs:run.font.size=Pt(9)
p('교육 항목별 완료 상태와 수업 후 메모를 함께 남겨 다음 수업에서 참고하는 방식입니다. 화면 속 과정명과 표시는 기능 설명을 위한 것이며, 교육단체의 공식 자격 발급을 대신하지 않습니다.')

page('3 성장전략')
h2('강사 구독과 무료 체험')
p('이용료는 강사에게만 월 3,900원을 받는 방식으로 계획하고 있습니다. 처음 3개월은 무료로 제공해 본인의 수업에 적용해 볼 시간을 주려 합니다. 수강생은 별도 구독료 없이 함께 사용하도록 할 계획입니다. 강사가 일정 등록부터 수업 후 기록까지 익숙하게 사용할 수 있어야 무료 기간이 끝난 뒤에도 구독을 유지할 것으로 보고 있습니다.')
p('낮은 가격만으로 사업을 유지하기는 어렵습니다. 무료 체험 중 얼마나 자주 수업을 운영하는지, 체험이 끝난 뒤 실제로 결제하는지, 고객지원에 어느 정도 시간이 드는지를 함께 살펴 가격과 운영 방식을 조정하겠습니다.')
h2('유튜브 협업을 통한 초기 홍보')
p('현재 유튜브 채널 먹보야 놀자와 협업하고 있으며, 채널을 통한 AllBlue 홍보를 준비하고 있습니다. 앱의 기능을 나열하기보다 강사가 수업을 준비하고 교육 기록을 남기는 모습을 보여주는 콘텐츠를 구상하고 있습니다. 관심을 보인 강사가 무료 체험을 시작하고 실제 수업에 사용하는 과정까지 연결하는 것이 목표입니다.')
p('홍보 이후에는 채널을 통해 들어온 강사가 가입에 그치는지, 수업을 만들고 다시 사용하는지 확인하겠습니다. 광고 게시 등 다른 업체와의 비즈니스도 검토하고 있습니다. 현재 광고 매출은 없으며, 이용자 규모와 사용 빈도를 확인한 뒤 구체적인 조건을 협의할 계획입니다.')
h2('가격에 따른 구독 수입')
table(['유료 강사 수 가정','월 결제액','가입 후 12개월 결제액'],[
 ['100명','390,000원','3,510,000원'],
 ['300명','1,170,000원','10,530,000원'],
 ['1,000명','3,900,000원','35,100,000원']],[5.4,5.6,6])
p('위 금액은 고객 확보 목표가 아닌 가격 계산 예시입니다. 전원이 3개월 무료 체험 후 9개월간 구독을 유지한다고 가정했습니다. 세금, 결제 수수료와 운영비를 차감하기 전 결제액이며 광고 수입은 포함하지 않았습니다.')
p('유료 강사가 300명이어도 월 결제액은 117만 원입니다. 초기에는 협업 채널을 활용해 고객을 확보하는 비용을 줄이고, 서버와 문서 보관 비용 및 고객지원 시간을 관리해야 합니다. 국내에서 반복 이용과 유료전환이 확인되면 여러 강사가 함께 운영하는 교육업체의 수요를 검토하겠습니다. 해외 진출은 국내 운영 경험을 쌓은 뒤 판단할 계획입니다.')

page('3 사업화 추진계획')
h2('현재 실적과 사업자 매출')
p('AllBlue는 운영 환경 테스트 단계이며 아직 매출이 없습니다. 강사와 수강생을 합한 사용자는 5명 이내입니다. 출시 후에는 소수의 강사가 수업 준비부터 기록까지 실제로 사용하는 과정을 먼저 살펴보겠습니다.')
table(['구분','금액 또는 상태','기준'],[
 ['AllBlue 자체 매출','없음','현재까지의 서비스 매출'],
 ['럼바우 2025년 매출','0원','대표자 확인 기준 사업자 매출'],
 ['럼바우 2026년 월 매출','6,700,000원','다른 프로젝트 매출, 부가세 제외'],
 ['럼바우 2026년 예상매출','80,400,000원','월 670만 원이 12개월 지속되는 가정']],[5,4.2,7.8])
p('럼바우의 기존 프로젝트 매출은 AllBlue의 사업 실적과 구분합니다. 2026년 금액은 1월부터 발생한 월 매출이 연말까지 이어진다는 전제로 산정한 예상치입니다. 2025년 개인 프리랜서 활동으로 얻은 수입은 위 사업자 매출에 포함하지 않았습니다.')
h2('출시 이후 진행할 일')
table(['시기','실행 내용','확인할 결과'],[
 ['출시 초기','운영 환경 점검과 마켓 공개, 협업 채널 홍보 준비','가입 후 첫 수업까지 진행하는 데 어려움이 없는지'],
 ['강사별 무료 체험 3개월','실제 교육과정에 적용하고 사용 의견 수집','수업을 다시 운영하는지, 어떤 기능에서 불편을 겪는지'],
 ['강사별 체험 종료 후','유료전환 안내와 이용 지속 여부 확인','실제 결제와 해지 사유, 강사당 지원 비용']],[3.6,6.7,6.7])
h2('개선 여부를 판단하는 기준')
p('활성 강사 수는 일정 기간에 수업을 만들고 운영한 강사 수로 확인하겠습니다. 완료한 수업 중 디브리핑을 남긴 비율, 첫 수업 이후 다시 앱을 사용한 비율도 함께 보겠습니다. 아직 측정 결과는 없으며 초기 사용 과정부터 기록을 쌓을 계획입니다.')
p('강사가 기록을 남기지 않는다면 입력 과정이 번거로운지, 기록할 필요를 느끼지 못하는지부터 살펴보겠습니다. 기능을 추가하기 전에 사용이 끊기는 이유를 파악하고, 반복 이용이 자리 잡은 뒤 고객 모집을 늘리겠습니다.')

page('4 팀 구성')
h2('개발과 프리다이빙 경험')
p('럼바우 대표 남상현입니다. 개발 14년차이며 프리다이빙은 3년째 하고 있습니다. AIDA Level 2 자격을 보유하고 있습니다. 현재 내부 인력은 저 한 명으로, AllBlue의 개발과 운영을 맡고 있습니다.')
p('AIDA Instructor 자격을 가진 협업자가 있고, 유튜브 먹보야 놀자와 홍보를 준비하고 있습니다. 개발 경험을 바탕으로 앱을 만들고, 강사의 의견과 실제 이용 결과를 반영해 수정하는 방식으로 제품을 다듬어 갈 계획입니다.')
h2('사업자 현황')
table(['항목','내용'],[
 ['상호와 대표자','럼바우 / 남상현'],
 ['사업 형태','개인사업자'],
 ['개업일','2025년 4월 7일'],
 ['사업장','서울특별시 구로구'],
 ['업태와 종목','정보통신업 / 컴퓨터 프로그래밍 서비스업'],
 ['내부 인력','대표자 1인']],[4,13])
h2('현재 인력으로 운영하는 방법')
p('개발과 고객 대응을 혼자 맡고 있어, 초기에는 수업 운영에 꼭 필요한 기능부터 안정화할 계획입니다. 사용 의견은 수업을 진행하기 어려운 오류와 편의 개선 요청으로 나누고, 수업에 영향을 주는 문제를 먼저 해결하겠습니다. 무료 체험 기간의 문의와 대응 시간을 기록해 이용자가 늘어날 때 필요한 운영 부담도 파악하겠습니다.')
doc.add_heading('5 기타 사항',1)
p('현재까지 투자유치, 수상 및 정부지원사업 수행 이력은 없습니다.')
p('이번 경진대회에서는 월 3,900원의 구독료로 사업을 유지하기 위한 고객 규모와 비용 구조를 점검받고 싶습니다. 초기 강사가 체험을 마친 뒤에도 서비스를 계속 사용할 이유를 구체화하고, 실제 이용 결과에 맞춰 사업모델을 보완하겠습니다.')

for part in [doc.styles.element,doc.element]:
 for border in list(part.iter(qn('w:pBdr'))): border.getparent().remove(border)
for style in doc.styles:
 if style.type==1:
  rf=style.element.get_or_add_rPr().get_or_add_rFonts()
  for attr in list(rf.attrib):
   if 'Theme' in attr or 'theme' in attr: del rf.attrib[attr]
  for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'SUIT')
# Respect the declared column widths in Word and LibreOffice.
for t in doc.tables:
 widths=[c.width for c in t.rows[0].cells]
 for col,w in zip(t.columns,widths):col.width=w
doc.core_properties.title='AllBlue G밸리 창업경진대회 사업계획서'
doc.core_properties.author=''
doc.save(ROOT/'AllBlue_Gvalley_business_plan_DRAFT.docx')
print(ROOT/'AllBlue_Gvalley_business_plan_DRAFT.docx')
