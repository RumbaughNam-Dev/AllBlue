from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json

ROOT=Path(__file__).resolve().parents[1]
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8)
sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 s=doc.styles[name];s.font.name='SUIT';s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'SUIT')
 s.font.size=Pt(10.5)
 s.paragraph_format.space_after=Pt(7)
 s.paragraph_format.line_spacing=1.25
doc.styles['Title'].font.size=Pt(21)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
h=sec.header.paragraphs[0];h.text='AllBlue  |  G밸리 창업경진대회  |  검토용 초안 · 제출 금지'
h.runs[0].font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('2026년 10월 7일 작성  ·  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
for r in f.runs:r.font.size=Pt(8)

def p(text,style=None):return doc.add_paragraph(text,style)
def h2(text):doc.add_heading(text,2)
def page(title):
 h=doc.add_heading(title,1);h.paragraph_format.page_break_before=True
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,txt in zip(t.rows[0].cells,headers):c.text=txt
 for row in rows:
  for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
 for ri,row in enumerate(t.rows):
  for ci,c in enumerate(row.cells):
   if widths:c.width=Cm(widths[ci])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   pr.append(b)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'95');x.set(qn('w:type'),'dxa');margins.append(x)
   pr.append(margins)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');pr.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.15
    for r in para.runs:r.font.size=Pt(9);r.bold=ri==0
  nr=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(nr)
 rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
 p('')
 return t

doc.add_heading('AllBlue 사업계획서 초안',0)
p('2026 제10회 G밸리 창업경진대회 서식 2 작성용')
p('검토 안내  |  현재 코드에서 확인한 기능과 사업화 제안을 구분했습니다. [확인 필요] 항목과 계획안을 대표자가 검토한 뒤 원본 신청서에 반영해야 합니다. 이 파일은 서명·증빙을 포함한 최종 제출본이 아닙니다.')
doc.add_heading('1 문제인식',1)
h2('창업아이템 개요')
p('AllBlue는 프리다이빙 강사와 교육생이 일정, 수업 참여, 교육 진도, 동의서 확인과 수업 후 피드백을 연결해 관리하는 모바일 서비스다. 강사가 한 차례의 수업을 준비하고 마무리하는 과정에서 필요한 정보를 참여자별로 확인하도록 구성했다. 초기 사업화 대상은 반복적으로 교육과정을 운영하는 프리다이빙 강사와 소규모 교육업체다.')
h2('개발배경과 해결할 문제')
p('검증할 고객 문제는 수업 정보가 일정표, 대화방, 문서와 개인 기록에 나뉘었을 때 발생하는 반복 확인 업무다. 강사는 참석 여부를 확인한 뒤 참가자의 서류 준비 상태와 이전 교육내용을 다시 찾아야 한다. 수업 이후에도 이수 항목과 다음 교육 목표를 기록해야 한다. 이 과정의 분산이 얼마나 큰 부담인지는 실제 강사 인터뷰와 사용 전후 측정으로 확인한다.')
p('PADI는 프리다이빙 교육을 단계별 과정으로 안내하고 있다[1]. 이는 교육 이력을 지속해서 관리할 사용 맥락의 근거이며, 국내 시장규모나 AllBlue의 수요가 이미 검증되었다는 근거로 사용하지 않는다.')
table(['구분','초기 대상','검증할 구매 이유'],[
 ['주요 구매자','프리다이빙 강사·소규모 교육업체','수업 준비와 교육기록 관리에 드는 시간 감소'],
 ['함께 사용하는 이용자','강사에게 교육받는 수강생','일정 확인과 교육 피드백 접근의 편의'],
 ['후속 확장 대상','해양 체험·다이빙센터','여러 강사·반복 방문객 관리의 필요성']],[2.4,6,8.6])
p('[확인 필요] 대표자가 이 문제를 직접 경험한 사례, 인터뷰 대상과 날짜, 현재 확보한 테스트 강사 수. 인터뷰 전에는 위 문제를 검증 가설로 제시한다.')

page('2 해결방안')
h2('수업 전후를 연결하는 사용 흐름')
p('강사는 일정과 참가자를 등록하고, 수강생은 초대에 응답한다. 강사는 일정 상세에서 참여자별 동의서 상태와 교육 관련 정보를 확인한다. 수업을 마친 뒤에는 이수 항목과 디브리핑 메모를 남긴다. 다음 수업에서는 누적된 진도와 피드백을 참조한다. 이 흐름은 현재 프런트엔드 코드에 존재하는 화면과 API 연결을 기준으로 정리했다. 운영 환경에서의 전체 동작과 고객 이용 실적은 별도로 확인해야 한다.')
table(['업무 단계','확인한 기능','확인 범위'],[
 ['수업 구성','일정 작성·수정, 수업 유형과 참가자 설정','일정 편집 화면·API'],
 ['참석 확인','초대 수락·거절, 알림 조회','초대 상태·알림 화면·API'],
 ['서류 준비','면책동의서·의료진술서 서명 상태 표시와 문서 연결','일정 상세 화면·데이터 항목'],
 ['교육 진행','교육 요구사항별 완료 상태 조회·변경','교육기록 화면·API'],
 ['수업 마무리','디브리핑 메모 저장·이력 조회','디브리핑 화면·API'],
 ['관계 관리','강사·교육생·친구 관리','관계 관리 화면']],[2.3,8.1,6.6])
h2('이용자가 얻을 가치와 측정 방법')
p('첫 번째 가치는 한 수업의 준비 상태를 한 곳에서 확인하는 편의다. 두 번째는 이전 교육내용을 다음 수업에 연결하는 연속성이다. 두 가치를 검증하기 위해 수업 준비에 걸리는 시간, 참여자 상태를 확인하기 위해 오가는 연락 횟수, 수업 후 기록 작성률을 측정한다. 아직 측정되지 않은 시간 절감률이나 안전사고 감소율을 성과로 기재하지 않는다.')
h2('제출 전 제품 증거 보강')
p('[확인 필요] 실제 앱에서 개인정보를 가린 일정 상세 화면, 교육진도 화면, 디브리핑 화면을 확보한다. 화면에 나타나는 기능과 설명을 일치시키고, 개발용 화면은 개발용이라고 표기한다. 현재 초안에는 실제 화면으로 오인할 수 있는 합성 이미지를 넣지 않았다.')

page('2 해결방안과 경쟁력')
h2('기술 구현과 운영 준비')
p('앱은 Expo와 React Native 기반으로 구성되어 있으며, 일정·참가자·교육 요구사항·디브리핑을 서버 API로 연결한다. 일정 초대는 대기·수락·거절 등 상태를 구분하고, 사용자 역할과 참가 상태에 따라 화면에서 제공하는 작업을 나눈다. 프런트엔드의 권한 표시는 서버의 접근권한 검사를 대신하지 않으므로 서버 동작 검증은 출시 점검 항목으로 둔다.')
p('현재 확인한 기술적 기반은 업무 데이터를 연결하는 앱 구조다. AI 학습모델, 등록특허, 독점 기술 또는 의료적 효과가 확보되었다고 주장하지 않는다. 상용화 경쟁력은 기능 존재 자체보다 현장 사용성, 데이터의 연속성, 반복 이용과 유료전환 실적으로 입증할 계획이다.')
h2('경쟁과 대체수단')
table(['대안','확인하거나 검증할 내용','AllBlue의 검증 방향'],[
 ['MySSI','공식 안내에 디지털 교육·자격증·로그북 기능 존재[2]','기록 기능만으로 차별화했다고 주장하지 않음'],
 ['메신저·일정표·문서 조합','강사가 실제 어떤 조합을 쓰는지 인터뷰 필요','한 수업의 참석·서류·진도·피드백을 연결하는 편의 비교'],
 ['AllBlue','관련 화면과 API 연결을 코드에서 확인','실제 사용성·반복 사용·비용 지불 의사 검증']],[3,7,7])
p('MySSI에 특정 기능이 없다는 비교는 하지 않는다. 동일한 사용자 과업을 실제로 비교한 뒤 차별점을 확정한다. 교육기관의 공식 인증 발급이나 제휴 권한을 보유한 것으로 표현하지 않는다.')
h2('고도화 우선순위')
p('첫째, 초대 응답부터 수업 후 기록까지 핵심 흐름을 실기기로 검증한다. 둘째, 강사·교육생 사이의 정보 접근 범위와 문서 조회 권한을 점검한다. 셋째, 실제 수업에서 반복 사용이 확인된 뒤 운영 지표와 유료 기능을 설계한다. 의료진술서 등 민감할 수 있는 정보의 수집 범위와 보관·삭제 절차는 전문가 검토를 받아 확정한다.')

page('3 성장전략')
h2('초기 고객 확보와 사업모델')
p('현재 강사와 수강생을 합한 실사용 규모는 5명 이내이며 운영 환경에서 테스트하고 있다. 유튜브 채널 먹보야 놀자와 협업 중이며, 해당 채널을 통한 홍보를 계획하고 있다. 수업 준비부터 교육기록까지 실제 사용 장면을 콘텐츠로 소개하고, 관심 강사가 교육과정에 도입하도록 유도한다. 채널의 구독자 수·예상 노출·계약 조건은 확인되지 않아 성과나 확정 유입으로 기재하지 않는다.')
p('과금 계획은 강사에게만 월 3,900원을 부과하고 최초 3개월은 무료로 제공하는 방식이다. 무료 기간에는 일정 생성 이후 교육기록까지 이어지는 반복 사용을 확인한다. 이용료 부과와 결제 기능의 실제 구현 상태는 별도 확인한다. 타 업체와의 광고 게시 등 비즈니스 활동을 추가 수익원으로 검토하되 광고 계약·단가·매출이 확보된 것으로 기재하지 않는다.')
h2('구독 수익 구조와 검증 과제')
table(['유료 전환 강사 수 가정','유료 기간 월 결제액','가입 후 12개월 결제액'],[
 ['100명','390,000원','3,510,000원'],
 ['300명','1,170,000원','10,530,000원'],
 ['1,000명','3,900,000원','35,100,000원']],[5.4,5.6,6])
p('위 표는 목표나 매출 예측이 아닌 가격 구조의 단순 계산이다. 전원이 무료 3개월 후 전환하고 이후 9개월간 해지하지 않는 가정이며, 부가세·결제 및 앱마켓 수수료·환불·운영비를 차감하지 않은 결제액이다. 실제 이용자 수 5명 이내와 위 가정치를 혼동하지 않는다. 광고 수입은 포함하지 않았다.')
p('낮은 이용료는 초기 도입 부담을 줄일 수 있지만 상당한 유료 강사 기반이 형성되기 전에는 구독 수입이 작다. 따라서 광고비 지출보다 협업 채널과 강사 추천 경로의 실제 전환을 먼저 측정하고, 고객지원·문서 보관·서버 비용을 반영해 고객당 손익을 검증한다.')
h2('목표시장과 확장 방향')
p('1차 시장은 반복 교육을 운영하는 국내 프리다이빙 강사다. 국내 유효 강사 수에 대한 확인된 통계가 없어 전체 시장규모를 임의로 제시하지 않는다. 실제 접촉 가능한 강사 목록, 월간 교육 빈도와 유료 전환 의향을 바탕으로 목표 고객 수를 산정한다. 대표자와 협업 채널이 도달할 수 있는 고객 규모는 추가 확인이 필요하다.')
p('국내 반복 사용이 확인되면 다수 강사 교육업체의 공동 운영 수요와 해외 강사의 업무 적합성을 순차적으로 검토한다. 해외진출 실적이나 교육단체의 공식 제휴는 현재 사업 성과로 기재하지 않는다.')

page('3 사업화 추진계획과 검증지표')
h2('현재까지의 준비와 확인할 실적')
p('Google Play는 검토 중이며 App Store는 검토가 완료되었다. 운영 환경 테스트 단계로 실제 사용자는 강사·수강생 합계 5명 이내다. 일정·교육기록·문서 상태 조회 화면과 API 연결을 코드에서 확인했다. AllBlue 매출은 현재 없으며, 럼바우의 다른 프로젝트 매출과 구분한다. 사업자 전체 매출액과 정확한 사용자 구분은 추가 확인한다.')
table(['기간','실행 계획안','결과물'],[
 ['출시 초기','마켓 공개 상태 확인, 운영 환경 점검, 협업 채널 소개 준비','공개 링크, 실제 사용 화면, 기능 점검표'],
 ['사용자별 1~3개월','무료 체험 중 교육과정 적용, 반복 사용과 개선 요구 수집','사용 이력, 재사용 지표, 개선 우선순위'],
 ['사용자별 4개월부터','무료 체험 종료 후 실제 유료전환과 이탈 원인 측정','결제 전환·해지·지원 비용, 가격 유지 여부']],[2.8,8.5,5.7])
p('위 일정은 대표자 검토 전 계획안이다. 실제 고객 모집 상황과 개발·운영 가능 시간에 맞춰 확정한다. 경진대회의 멘토링과 데모데이 참여 일정도 함께 반영해야 한다.')
h2('검증할 지표')
table(['지표','산정 기준','현재 값'],[
 ['전체 실사용자','강사·수강생 합계, 정확한 구분은 추가 확인','5명 이내'],
 ['활성 강사','기간 내 수업을 생성하고 운영한 고유 강사 수','구분 확인 필요'],
 ['수업 후 기록률','디브리핑 기록이 남은 완료 수업 / 완료 수업','확인 필요'],
 ['반복 이용','첫 수업 이후 다시 수업을 운영한 강사 비율','확인 필요'],
 ['업무시간 변화','동일 업무의 도입 전후 소요시간 비교','측정 전'],
 ['유료전환','유료 제안 대상 중 실제 결제한 고객 비율','확인 필요']],[3.2,10.4,3.4])
h2('실행 위험과 대응')
p('반복 사용이 낮으면 유료화보다 도입 과정과 필수 기능을 먼저 개선한다. 지원금이나 상금을 전제로 고정비를 늘리지 않는다. 상금은 입상 결과에 따라 달라지며 확보된 자금으로 표시하지 않는다. 운영비·고객지원·문서 저장 비용을 확인한 뒤 사업 유지에 필요한 비용을 산정한다.')

page('4 팀 구성과 기타 사항')
h2('대표자와 참여인력')
p('럼바우는 2025년 4월 7일 개업한 서울 구로구 소재 개인사업자이며, 업태·종목은 정보통신업·컴퓨터 프로그래밍 서비스업이다. 대표자 남상현은 개발 14년차이며 프리다이빙 3년차로 AIDA Level 2 자격을 보유하고 있다. 동업자는 AIDA Instructor 자격을 보유한다. 아래 경력과 자격은 대표자 제공 정보이며 제출 전 증빙을 대조한다.')
p('내부 운영 인력은 대표자 1인이며 유튜브 먹보야 놀자와 외부 협업하고 있다. 대표자의 개발 경험과 프리다이빙 경험, 강사 자격을 가진 협업자의 전문성을 제품과 현장 검증에 연결할 수 있다는 점이 팀의 강점이다. 협업자를 직원 수에 포함하거나 공동대표로 기재하지 않는다. 담당 업무와 참여 범위는 확인 후 확정한다.')
table(['항목','작성할 내용'],[
 ['사업자 상호·대표자','럼바우 / 남상현 · 사업자등록증명 기재내용 확인'],
 ['개발 전문성','개발 14년차. 세부 경력 연도·대표 프로젝트·담당 역할 추가 확인'],
 ['프리다이빙 전문성','프리다이빙 3년차, AIDA Level 2. 자격증과 현장 사례 확인'],
 ['참여인력','내부 인력 대표자 1인. AIDA Instructor 동업자와 협업 범위 확인'],
 ['외부 협력','유튜브 먹보야 놀자와 협업 중, 채널 홍보 예정. 링크·역할 확인']],[4.3,12.7])
h2('5 기타 사항')
table(['구분','현재 기재 상태'],[
 ['지식재산권 등록','미확인. 등록 완료한 권리만 확인 후 기재'],
 ['수상경력','없음 · 대표자 확인'],
 ['정부지원사업 수행','없음 · 대표자 확인'],
 ['투자유치','없음 · 대표자 확인'],
 ['제1~9회 G밸리 수상 여부','수상 이력 없음 · 대표자 확인']],[4.3,12.7])
page('작성 근거와 제출 준비')
h2('참고자료와 근거의 범위')
p('[1] PADI, Freediving Program Information. 프리다이빙 교육과정 소개. https://www.padi.com/education/freediving (2026년 10월 7일 확인)')
p('[2] SSI, MySSI App. 디지털 교육·자격증·로그북 기능 소개. https://www.divessi.com/en/join-ssi/register-free/myssi-app (2026년 10월 7일 확인)')
p('제품 근거는 2026년 10월 7일 확인한 AllBlue 저장소의 README, 일정 편집·상세 화면, 디브리핑·교육기록 화면과 API 정의다. 코드 존재는 확인했으나 이번 작업에서 실서비스를 실행하거나 성능·보안·사용성과를 검증하지 않았다.')
p('제출 전 점검  |  미확인 항목 보완, 대표자 경력·성과 증빙 대조, 실제 제품 화면 추가, 계획안 승인, 원본 서식 반영과 최종 PDF 검수가 필요하다. 검토용 안내와 미확인 표시를 남긴 채 제출하지 않는다.')

h2('최종 제출본에 필요한 자료')
table(['구분','준비할 자료'],[
 ['사업자 확인','사업자등록증명과 인적사항 확보. 공고 첨부명인 사업자등록증과의 대체 인정 여부 확인'],
 ['재무 증빙','2025년 신고는 세무사에게 위임. 제출 재무제표·증명 발급 여부 확인 후 미작성 시 대체 증빙 인정 여부 확인'],
 ['제품 증거','개인정보를 가린 실제 앱 화면, 공개된 경우 앱마켓 링크'],
 ['팀 증거','대표 개발 프로젝트·담당 역할, 자격 증빙, 협업 채널 링크'],
 ['서명 서류','원본 개인정보 동의서와 서약서의 내용 확인 및 대표자 서명']],[4.3,12.7])
p('이 페이지는 초안 검토를 위한 부록이다. 최종 제출 시 검토 안내를 제외하고 공고와 접수 화면의 파일 구성에 맞춰 정리한다.')

for part in [doc.styles.element,doc.element]:
 for border in list(part.iter(qn('w:pBdr'))): border.getparent().remove(border)
for style in doc.styles:
 if style.type==1:
  rf=style.element.get_or_add_rPr().get_or_add_rFonts()
  for attr in list(rf.attrib):
   if 'Theme' in attr or 'theme' in attr: del rf.attrib[attr]
  for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'SUIT')
# Respect the declared column widths in Word and LibreOffice.
for t in doc.tables:
 widths=[c.width for c in t.rows[0].cells]
 for col,w in zip(t.columns,widths):col.width=w
doc.core_properties.title='AllBlue G밸리 창업경진대회 사업계획서 검토용 초안'
doc.core_properties.author=''
doc.save(ROOT/'AllBlue_Gvalley_business_plan_DRAFT.docx')
print(ROOT/'AllBlue_Gvalley_business_plan_DRAFT.docx')
